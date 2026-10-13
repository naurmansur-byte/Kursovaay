"""Пакет для преподавателя: программа, диагностика, недели 1–2, инструкция и лист замечаний — по порядку и без служебных пометок.

Запуск: .venv/Scripts/python tools/teacher_pack.py [папка-результат]
Исходники не меняются: чистятся только копии (статусы «ждёт проверки…», ссылки на внутренние файлы, заметки о разработке).
Презентации берутся готовыми из exports/week-NN/lesson-K-slides.pdf (собрать: tools/slides/lesson_deck.py), памятка — exports/week-NN/memo.pdf.
"""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'tools'), str(ROOT / 'tools' / 'handout')]
from export_lesson_kit import export_docx  # noqa: E402
import handout  # noqa: E402

PACK = ROOT / 'course' / 'teacher-pack'
WEEKS = {'week-01': 'Неделя 1 — язык химии', 'week-02': 'Неделя 2 — формулы и уравнения'}

# Предложения с этими словами — заметки о разработке, преподавателю не нужны.
DEV = re.compile(r'\bМакс(?:а|у|ом|е)?\b|Claude|\bагент|Google|\bДиск|загрузк|в сессии|outline|EPUB|фактическ\w* чтени|принятие|ROADMAP|STYLE|docs/|prompts/|'
                 r'\.md\b|\bT\d\d\b|экспортёр|PPTX|не создан|CONNECT|разработк|оглавлени', re.I)
REPLACE = [
    (r'\*\*(?:Статус: )?[Чч]ерновик[^*]*\*\*', '**Черновик для проверки преподавателем.**'),
    (r'\*\*Только преподавателю;[^*]*\*\*', '**Только преподавателю.**'),
    (r'\[([^\]]+)\]\((?!https?:)[^)]*\)', r'\1'),            # внутренние ссылки → текст
    (r'диагностик(\w+) T02', r'входной диагностик\1'), (r'полного T02', 'полной входной диагностики'),
    (r'с T02\b', 'с входной диагностикой'), (r'по T02\b', 'по входной диагностике'), (r'\bT02\b', 'входная диагностика'),
    (r'официальн\w+ EPUB', 'учебника'),
]
SENT = re.compile(r'(?<=[.!?…])\s+(?=[А-ЯЁA-Z«*(])')
DROPPED: list[str] = []  # что вырезано — для проверки глазами (пишется в exports/teacher-pack-dropped.txt)


DEV_SECTION = re.compile(r'^(#{2,4}) .*(Google|Диск|Перенос в|Передача|Что проверено|разработк)', re.M)


def clean(text: str) -> str:
    # Целые разделы о переносе на Диск и о ходе разработки — убрать до следующего заголовка того же уровня.
    while m := DEV_SECTION.search(text):
        nxt = re.compile(rf'^#{{1,{len(m.group(1))}}} ', re.M).search(text, m.end())
        DROPPED.append('[раздел] ' + text[m.start():m.end()])
        text = text[:m.start()] + (text[nxt.start():] if nxt else '')
    out = []
    for line in text.splitlines():
        for pat, rep in REPLACE:
            line = re.sub(pat, rep, line)
        if line.lstrip().startswith('|') or not DEV.search(line):
            out.append(line)
            continue
        kept = [s for s in SENT.split(line) if not DEV.search(s)]   # убрать только «служебные» предложения
        DROPPED.extend(s for s in SENT.split(line) if DEV.search(s))
        if kept and ''.join(kept).strip(' -*>'):
            out.append(' '.join(kept))
    return re.sub(r'\n{3,}', '\n\n', '\n'.join(out))


def section(md: str, title: str) -> str:
    m = re.search(rf'^## {re.escape(title)}[^\n]*\n(.*?)(?=^## |\Z)', md, re.S | re.M)
    return m.group(1) if m else ''


def program() -> str:
    src = (ROOT / 'course' / 'PROGRAM.md').read_text(encoding='utf-8')
    grid = '\n'.join(ln for ln in section(src, 'Календарная сетка').splitlines() if ln.startswith('|'))
    rows = []
    for ln in grid.splitlines():   # строки пауз — коротко, без служебных пояснений
        cells = [c.strip() for c in ln.strip('|').split('|')]
        if len(cells) > 2 and cells[1] == 'пауза':
            kind = 'Зимняя' if '.01' in cells[0] or '.12' in cells[0] else 'Весенняя'
            ln = f'| {cells[0]} | пауза | {kind} пауза (школьные каникулы) | — | — |'
        rows.append(ln)
    cover = section(src, 'Покрытие 19 номеров спецификации')
    cover_table = '\n'.join(ln for ln in cover.splitlines() if ln.startswith('|'))
    note = next((ln for ln in cover.splitlines() if ln.startswith('В №07')), '').replace('В №07 PDF действительно печатает', 'В №07 спецификация печатает')
    detail = '\n'.join(ln for ln in section(src, 'Первые три содержательные недели подробно').splitlines() if ln.startswith('|'))
    detail = detail.replace('Внутренние ориентиры', 'Учебник 8 класса').replace('; содержание этих разделов прочитать до разработки У3', '') \
        .replace('; по проверенному оглавлению', '').replace(' по оглавлению', '')
    return ((PACK / 'program-intro.md').read_text(encoding='utf-8')
            + '\n## Календарь по неделям\n\nК — календарная неделя, У — учебная неделя (и номера уроков).\n\n' + '\n'.join(rows)
            + '\n\n## Покрытие спецификации ЕНТ «Химия» (с 2026 года)\n\nУ — учебная неделя курса.\n\n' + cover_table + '\n\n' + note
            + '\n\n## Недели 1–3 подробно\n\nУчебник: Оспанова М. К., Аухадиева К. С., Белоусова Т. Г. «Химия. 8 класс», Мектеп, 2018.\n\n' + detail + '\n')


def feedback_form(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(1.5)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading('Лист замечаний', level=1)
    doc.add_paragraph('Номер файла — из названия (например, «03-05» — неделя 1, файл 05). Важность: К — критично (ошибка в химии или ответе), Ж — желательно.')
    head = ['№', 'Файл', 'Страница / слайд / задание', 'Что не так', 'Как правильно', 'К / Ж']
    widths = [Cm(1), Cm(2.2), Cm(4.3), Cm(8.5), Cm(8.5), Cm(1.5)]
    table = doc.add_table(rows=1, cols=len(head))
    table.style = 'Table Grid'
    for cell, text, w in zip(table.rows[0].cells, head, widths):
        cell.text, cell.width = text, w
        cell.paragraphs[0].runs[0].bold = True
    for k in range(1, 26):
        cells = table.add_row().cells
        cells[0].text = str(k)
        for cell, w in zip(cells, widths):
            cell.width = w
    doc.add_paragraph()
    doc.add_paragraph('Общее впечатление и вопросы:')
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def docx(md_text: str, path: Path) -> None:
    export_docx(clean(md_text), path)


def build(out: Path) -> Path:
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    handout.main(str(PACK / '00-start.md'), str(out / '00 Начните отсюда.pdf'))
    (out / '00 Начните отсюда.html').unlink()
    docx(program(), out / '01 Программа курса.docx')
    diag = ROOT / 'course' / 'diagnostic'
    docx((diag / 'student.md').read_text(encoding='utf-8'), out / '02 Входная диагностика' / '1 Лист ученика.docx')
    docx((diag / 'teacher.md').read_text(encoding='utf-8'), out / '02 Входная диагностика' / '2 Ключ и распределение по группам.docx')
    for n, (week, title) in enumerate(WEEKS.items(), 3):
        src, exp, dst = ROOT / 'course' / week, ROOT / 'exports' / week, out / f'0{n} {title}'
        dst.mkdir(parents=True)
        md = lambda name: (src / f'{name}.md').read_text(encoding='utf-8')  # noqa: E731
        student = {}
        for name in ['notes-student', 'pre-work', 'homework', 'quiz']:
            pdf = exp / 'student' / f'{name}.pdf'
            handout.main(str((src / f'{name}.md').relative_to(ROOT)), str(pdf.relative_to(ROOT)))
            student[name] = pdf
        docx(md('lesson-1-plan'), dst / '01 Урок 1 — план.docx')
        shutil.copy(exp / 'lesson-1-slides.pdf', dst / '02 Урок 1 — презентация.pdf')
        shutil.copy(student['notes-student'], dst / '03 После урока 1 — конспект.pdf')
        shutil.copy(student['pre-work'], dst / '04 После урока 1 — закрепление.pdf')
        shutil.copy(student['homework'], dst / '05 Домашнее задание.pdf')
        docx(md('homework-solutions'), dst / '06 Домашнее задание — решения и критерии.docx')
        docx(md('lesson-2-plan'), dst / '07 Урок 2 — план.docx')
        shutil.copy(exp / 'lesson-2-slides.pdf', dst / '08 Урок 2 — презентация.pdf')
        shutil.copy(student['quiz'], dst / '09 Контроль — лист ученика.pdf')
        docx(md('quiz'), dst / '10 Контроль — с ключом.docx')
        shutil.copy(exp / 'memo.pdf', dst / '11 Памятка недели.pdf')
        docx(md('teacher-notes'), dst / '12 Заметки преподавателю.docx')
    feedback_form(out / '05 Лист замечаний.docx')
    archive = out.with_suffix('.zip')
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in sorted(out.rglob('*')):
            if f.is_file():
                z.write(f, f.relative_to(out.parent))
    return archive


if __name__ == '__main__':
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'exports' / 'Для преподавателя — проверка курса'
    zip_path = build(target.resolve())
    (ROOT / 'exports' / 'teacher-pack-dropped.txt').write_text('\n\n'.join(DROPPED), encoding='utf-8')
    print(f'Teacher pack -> {target}\nZIP -> {zip_path} ({zip_path.stat().st_size / 1e6:.1f} MB)')
