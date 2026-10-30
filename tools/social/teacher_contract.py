"""Договор с преподавателем: шаблон с условиями курса и местами для подписи — PDF для печати и DOCX для правки.

Запуск из корня: .venv/Scripts/python.exe -I tools/social/teacher_contract.py [биологии|химии]
Результат: «00 Материалы ХимБио просто/Найм преподавателей/Договор с преподавателем <предмета>.pdf» и .docx.
Пустые поля (____) заполняются от руки или в Word; персональные данные в репозиторий не вносить.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools' / 'handout'))
from handout import chrome  # noqa: E402

FONTS = (ROOT / 'content/brand/fonts').as_uri()
OUT = ROOT / '00 Материалы ХимБио просто' / 'Найм преподавателей'
B = '________________'
S = '________'


def sections(subj: str) -> list:
    return [
        ('Предмет договора', [
            f'Исполнитель оказывает услуги преподавателя {subj} на онлайн-курсе подготовки к ЕНТ «ХимБио просто» (далее — курс): проводит уроки, проверяет домашние задания учеников и учебные материалы курса.',
            'Уроки проходят онлайн в Google Meet, домашние задания и вопросы учеников — в Telegram. Расписание групп — в приложении 1.',
        ]),
        ('Срок', [
            f'Договор действует с {S} 20__ г. по {S} 20__ г.',
            f'Первая неделя работы — испытательная: {S} (да / нет). По её итогам любая сторона может прекратить договор без срока предупреждения из раздела 7.',
        ]),
        ('Обязанности Исполнителя', [
            'Проводить два урока по 90 минут в неделю в каждой назначенной группе по расписанию из приложения 1.',
            'Вести урок по комплекту недели курса: план, презентация, конспект, домашнее задание, контроль. Не заменять материалы курса своими без согласования с Заказчиком.',
            f'Проверять комплект недели заранее — не позднее чем за {S} дней до первого урока недели — и присылать замечания Заказчику. Ошибку в фактах или ключах сообщать сразу.',
            f'Проверять домашние задания учеников в Telegram в течение {S} часов после их получения; назначать ученику вариант A, B или C и вести карту пробелов.',
            'Проводить контроль на уроках и учебные пробники в конце курса, разбирать их результаты с учениками.',
            'Объяснять по учебникам из перечня НЦТ и по правилам курса: научно верно, ошибки учебников — с пояснением «в учебнике написано…, верно…».',
            'Не обещать ученикам и родителям баллы ЕНТ, грант или поступление.',
            f'О переносе или отмене урока сообщать Заказчику не позднее чем за {S} часов. Замену преподавателя согласует Заказчик.',
            'Вести урок с исправной связью, камерой и микрофоном, из тихого места. Записывать уроки по правилам курса.',
            'Общаться с учениками вежливо; вопросы родителей об оплате, расписании и результатах передавать Заказчику.',
        ]),
        ('Обязанности Заказчика', [
            'Предоставить материалы курса: комплекты недель, презентации, ключи и решения, программу.',
            'Составить расписание групп, дать доступ к группам Telegram и встречам Google Meet.',
            'Набирать учеников и вести расчёты с учениками и родителями.',
            'Выплачивать вознаграждение в размере и сроки из раздела 5.',
            f'Сообщать Исполнителю об изменении расписания не позднее чем за {S} часов.',
        ]),
        ('Вознаграждение', [
            f'Вознаграждение Исполнителя — {B} тенге за {B} (один урок 90 минут / месяц работы / другое).',
            f'Выплата — {S} раз(а) в месяц, не позднее {S} числа, способом: {B}.',
            'Урок, отменённый или не проведённый по вине Исполнителя, не оплачивается. Урок, отменённый Заказчиком позднее срока из пункта 4.5, оплачивается.',
            'Налоги и обязательные платежи каждая сторона уплачивает в соответствии со своим правовым статусом.',
        ]),
        ('Материалы курса и данные учеников', [
            'Материалы курса — конспекты, презентации, задания, ключи, решения, пробники, памятки, видео — принадлежат Заказчику. Исполнитель не передаёт их третьим лицам, не публикует и не использует на других курсах и занятиях — во время действия договора и после его окончания.',
            'Персональные данные учеников и родителей — имена, телефоны, результаты, переписка — Исполнитель использует только для работы на курсе, не передаёт третьим лицам и удаляет у себя после окончания договора.',
            'Записи уроков хранятся и передаются ученикам только через каналы курса.',
        ]),
        ('Прекращение договора', [
            f'Любая сторона может прекратить договор, письменно предупредив другую сторону не позднее чем за {S} дней.',
            'При прекращении договора Исполнитель передаёт Заказчику карту пробелов и результаты учеников, а Заказчик выплачивает вознаграждение за оказанные услуги.',
            'Споры стороны решают переговорами, а если не договорились — в порядке, установленном законодательством Республики Казахстан.',
            'Договор составлен в двух экземплярах, по одному для каждой стороны.',
        ]),
    ]


PARTY_ROWS = ['ФИО / наименование', 'ИИН / БИН', 'Удостоверение личности: №, кем и когда выдано', 'Адрес', 'Телефон', 'Банковские реквизиты (IBAN, банк)']

RULES = [
    'Порядок объяснения: почему → алгоритм → разные примеры → ловушка → ученик решает сам. Ключ — после попытки.',
    'Урок 2 недели: 20 минут вопросов и повторения, разбор домашнего задания, практика в формате ЕНТ, 10-минутный контроль.',
    'Факты, определения и числа — по учебникам из перечня НЦТ. Ошибку учебника не переносим: говорим научно верно.',
    'С учеником — на «ты», коротко, без лишних слов. Ошибку разбираем по причине, а не заменяем готовым ответом.',
    'Вариант A, B или C — маршрут обучения, а не прогноз балла ЕНТ.',
]


def subj_word(arg: str) -> str:
    return {'биологии': 'биологии', 'химии': 'химии'}[arg]


# ---------- PDF ----------

CSS = """
:root { --ink: #141714; --muted: #5C625B; --line: #CFD3C8; --accent: #1F7A4A;
  --display: 'Unbounded', 'Arial Black', sans-serif; --body: 'Manrope', 'Segoe UI', Arial, sans-serif; }
@page { size: A4; margin: 16mm 16mm 18mm; }
* { box-sizing: border-box; }
body { margin: 0; color: var(--ink); background: #fff; font: 500 11.5px/1.55 var(--body); }
.brand { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 2.5px solid var(--accent); padding-bottom: 6px; margin-bottom: 18px; }
.brand b { font: 800 12px/1 var(--display); color: var(--accent); }
.brand span { color: var(--muted); font-size: 10.5px; }
h1 { font: 800 19px/1.2 var(--display); margin: 0 0 6px; text-align: center; }
.sub { text-align: center; color: var(--muted); margin: 0 0 14px; }
.meta { display: flex; justify-content: space-between; margin-bottom: 12px; }
h2 { font: 800 12.5px/1.3 var(--body); margin: 16px 0 6px; text-transform: uppercase; letter-spacing: .06em; color: var(--accent); break-after: avoid; }
p { margin: 0 0 5px; text-align: justify; }
.cl { display: grid; grid-template-columns: 30px minmax(0, 1fr); }
.cl b { font-variant-numeric: tabular-nums; }
.sig { display: grid; grid-template-columns: 1fr 1fr; gap: 22px; margin-top: 12px; break-inside: avoid; }
.sig h3 { font: 800 12px/1.3 var(--body); margin: 0 0 6px; }
.row { border-bottom: 1px solid var(--line); padding: 14px 0 2px; font-size: 9.5px; color: var(--muted); }
.signline { margin-top: 18px; display: flex; gap: 10px; align-items: flex-end; font-size: 10.5px; }
.signline span { flex: 1; border-bottom: 1px solid var(--ink); height: 18px; }
.app { break-before: page; }
table { width: 100%; border-collapse: collapse; margin-top: 8px; }
th, td { border: 1px solid var(--line); padding: 8px 8px 16px; text-align: left; vertical-align: top; font-size: 10.5px; }
th { background: #EEF3EE; padding-bottom: 8px; }
ol { margin: 4px 0 0; padding-left: 18px; } ol li { margin-bottom: 4px; }
"""


def clauses_html(secs: list) -> str:
    out = []
    for i, (title, items) in enumerate(secs, 1):
        out.append(f'<h2>{i}. {escape(title)}</h2>')
        out += [f'<div class="cl"><b>{i}.{j}.</b><p>{escape(t)}</p></div>' for j, t in enumerate(items, 1)]
    return '\n'.join(out)


def party_html(title: str) -> str:
    rows = ''.join(f'<div class="row">{escape(r)}</div>' for r in PARTY_ROWS)
    return f'<div><h3>{title}</h3>{rows}<div class="signline">Подпись <span></span> / <span></span></div></div>'


def html(subj: str, fonts: str) -> str:
    secs = sections(subj)
    sched_rows = ''.join('<tr><td></td><td></td><td></td><td></td><td></td></tr>' for _ in range(4))
    rules = ''.join(f'<li>{escape(r)}</li>' for r in RULES)
    return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">{fonts}<style>{CSS}</style></head><body>
<div class="brand"><b>ХимБио просто</b><span>онлайн-подготовка к ЕНТ по химии и биологии</span></div>
<h1>Договор возмездного оказания услуг № {S}</h1>
<p class="sub">преподаватель {subj} онлайн-курса «ХимБио просто»</p>
<div class="meta"><span>г. {B}</span><span>«____» {B} 20__ г.</span></div>
<p>{B}, далее — <b>Заказчик</b>, с одной стороны, и {B}, далее — <b>Исполнитель</b>, с другой стороны, вместе — стороны, заключили настоящий договор о следующем.</p>
{clauses_html(secs)}
<h2>{len(secs) + 1}. Реквизиты и подписи сторон</h2>
<div class="sig">{party_html('Заказчик')}{party_html('Исполнитель')}</div>
<section class="app">
<h2>Приложение 1. Расписание групп</h2>
<p>Приложение — часть договора. При изменении расписания стороны подписывают новое приложение.</p>
<table><tr><th>Группа</th><th>Дни недели</th><th>Время (Астана)</th><th>Первый урок</th><th>Число учеников</th></tr>{sched_rows}</table>
<div class="signline" style="margin-top:22px">Подпись Заказчика <span></span> Подпись Исполнителя <span></span></div>
<h2>Приложение 2. Правила преподавания на курсе</h2>
<p>Исполнитель ознакомлен с правилами и соблюдает их на уроках и в переписке с учениками.</p>
<ol>{rules}</ol>
<div class="signline" style="margin-top:28px">Ознакомлен(а), подпись Исполнителя <span></span> «____» {S} 20__ г.</div>
<div class="signline">Подпись Заказчика <span></span></div>
</section>
</body></html>"""


# ---------- DOCX ----------

def docx(subj: str, out: Path) -> None:
    d = Document()
    st = d.styles['Normal']
    st.font.name, st.font.size = 'Arial', Pt(10.5)
    green = RGBColor(0x1F, 0x7A, 0x4A)

    def para(text: str, bold: bool = False, align=None, size=None, color=None, space=4):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, p.paragraph_format.space_after = bold, Pt(space)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if align is not None: p.alignment = align
        return p

    para('ХимБио просто — онлайн-подготовка к ЕНТ по химии и биологии', bold=True, color=green, size=9)
    para(f'Договор возмездного оказания услуг № {S}', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    para(f'преподаватель {subj} онлайн-курса «ХимБио просто»', align=WD_ALIGN_PARAGRAPH.CENTER, space=10)
    para(f'г. {B}\t\t\t\t\t«____» {B} 20__ г.', space=10)
    para(f'{B}, далее — Заказчик, с одной стороны, и {B}, далее — Исполнитель, с другой стороны, вместе — стороны, заключили настоящий договор о следующем.')
    secs = sections(subj)
    for i, (title, items) in enumerate(secs, 1):
        para(f'{i}. {title}', bold=True, color=green, size=11, space=3).paragraph_format.space_before = Pt(8)
        for j, t in enumerate(items, 1):
            para(f'{i}.{j}. {t}', align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(f'{len(secs) + 1}. Реквизиты и подписи сторон', bold=True, color=green, size=11).paragraph_format.space_before = Pt(8)
    t = d.add_table(rows=len(PARTY_ROWS) + 2, cols=2)
    t.style = 'Table Grid'
    t.cell(0, 0).text, t.cell(0, 1).text = 'Заказчик', 'Исполнитель'
    for k, r in enumerate(PARTY_ROWS, 1):
        t.cell(k, 0).text = t.cell(k, 1).text = f'{r}:\n'
    t.cell(len(PARTY_ROWS) + 1, 0).text = t.cell(len(PARTY_ROWS) + 1, 1).text = 'Подпись ____________ / ____________'
    d.add_page_break()
    para('Приложение 1. Расписание групп', bold=True, color=green, size=11)
    para('Приложение — часть договора. При изменении расписания стороны подписывают новое приложение.')
    s = d.add_table(rows=5, cols=5)
    s.style = 'Table Grid'
    for k, h in enumerate(['Группа', 'Дни недели', 'Время (Астана)', 'Первый урок', 'Число учеников']):
        s.cell(0, k).text = h
    para('')
    para('Подпись Заказчика ____________     Подпись Исполнителя ____________', space=14)
    para('Приложение 2. Правила преподавания на курсе', bold=True, color=green, size=11)
    para('Исполнитель ознакомлен с правилами и соблюдает их на уроках и в переписке с учениками.')
    for k, r in enumerate(RULES, 1):
        para(f'{k}. {r}')
    para('')
    para('Ознакомлен(а), подпись Исполнителя ____________     «____» ________ 20__ г.')
    para('Подпись Заказчика ____________')
    d.save(out)


def main() -> None:
    subj = subj_word(sys.argv[1] if len(sys.argv) > 1 else 'биологии')
    local = (f"<style>@font-face {{ font-family: 'Unbounded'; font-weight: 200 900; src: url('{FONTS}/unbounded-cyrillic-wght-normal.woff2') format('woff2'); unicode-range: U+0400-045F; }}"
             f"@font-face {{ font-family: 'Unbounded'; font-weight: 200 900; src: url('{FONTS}/unbounded-latin-wght-normal.woff2') format('woff2'); }}"
             f"@font-face {{ font-family: 'Manrope'; font-weight: 200 800; src: url('{FONTS}/manrope-cyrillic-wght-normal.woff2') format('woff2'); unicode-range: U+0400-045F; }}"
             f"@font-face {{ font-family: 'Manrope'; font-weight: 200 800; src: url('{FONTS}/manrope-latin-wght-normal.woff2') format('woff2'); }}</style>")
    OUT.mkdir(parents=True, exist_ok=True)
    name = f'Договор с преподавателем {subj}'
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / 'contract.html'
        src.write_text(html(subj, local), encoding='utf-8')
        subprocess.run([chrome(), '--headless=new', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={OUT / (name + ".pdf")}', src.as_uri()],
                       check=True, capture_output=True)
    docx(subj, OUT / (name + '.docx'))
    print(OUT / name)


if __name__ == '__main__':
    main()
