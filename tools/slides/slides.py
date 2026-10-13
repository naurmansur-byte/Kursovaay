"""Презентация недели «ХимБио просто»: slides-outline.md -> слайды-картинки 1920×1080 (Chrome) -> PPTX с заметками + PDF.

Выглядит одинаково в PowerPoint любой версии, Google Slides и Meet: шрифты и индексы уже «впечатаны» в картинку.
Правки — в slides-outline.md, затем повторный запуск.
Запуск: .venv/Scripts/python tools/slides/slides.py week-01   (результат: exports/week-01/slides.pptx и slides.pdf)
"""
from __future__ import annotations

import html
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'tools' / 'handout'))
from handout import chrome  # noqa: E402

SUB = str.maketrans({'₀': '0', '₁': '1', '₂': '2', '₃': '3', '₄': '4', '₅': '5', '₆': '6', '₇': '7', '₈': '8', '₉': '9', 'ₘ': 'm', 'ₐ': 'A', 'ᵣ': 'r'})
SUP = str.maketrans({'⁰': '0', '¹': '1', '²': '2', '³': '3', '⁴': '4', '⁵': '5', '⁶': '6', '⁷': '7', '⁸': '8', '⁹': '9', '⁺': '+', '⁻': '−'})
CALLOUT = re.compile(r'^(Ловушка|Запомни|Задание[^:]*|Вопрос[^:]*|Входной вопрос|Проверь себя|Попробуй)\s*:\s*(.+)$', re.I)


def inline(text: str) -> str:
    """Markdown-строка -> HTML: **жирное**, `код`, ссылки, Unicode-индексы -> <sub>/<sup>."""
    text = html.escape(re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('`', ''))
    text = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', text)
    text = re.sub(r'[₀-₉ₘₐᵣ]+', lambda m: f'<sub>{m.group(0).translate(SUB)}</sub>', text)
    text = re.sub(r'[⁰-⁹⁺⁻]+', lambda m: f'<sup>{m.group(0).translate(SUP)}</sup>', text)
    return text


def body_html(lines: list[str], dark: bool) -> tuple[str, str]:
    """Пункты/нумерация/абзацы/таблицы -> HTML основного блока и выносок снизу."""
    out, callouts, table, listing = [], [], [], None

    def flush_table():
        nonlocal table
        if table:
            head, *rest = table
            out.append('<table><tr>' + ''.join(f'<th>{inline(c)}</th>' for c in head) + '</tr>'
                       + ''.join('<tr>' + ''.join(f'<td>{inline(c)}</td>' for c in r) + '</tr>' for r in rest) + '</table>')
            table = []

    def close_list():
        nonlocal listing
        if listing:
            out.append(f'</{listing}>')
            listing = None

    for line in lines:
        if line.startswith('|'):
            close_list()
            cells = [c.strip() for c in line.strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                table.append(cells)
            continue
        flush_table()
        m = CALLOUT.match(line.lstrip('- ').strip())
        if m and not dark:
            close_list()
            kind = 'trap' if m.group(1).lower().startswith('ловушка') else 'task'
            callouts.append(f'<div class="callout {kind}"><b>{html.escape(m.group(1))}</b>{inline(m.group(2))}</div>')
            continue
        item = re.match(r'^(?:([-*])|(\d+)[.)])\s+(.*)$', line)
        if item:
            tag = 'ol' if item.group(2) else 'ul'
            if listing != tag:
                close_list()
                out.append(f'<{tag}>')
                listing = tag
            out.append(f'<li>{inline(item.group(3))}</li>')
        else:
            close_list()
            out.append(f'<p>{inline(line)}</p>')
    close_list()
    flush_table()
    return '\n'.join(out), '\n'.join(callouts)


# Если текст не помещается в слайд, уменьшить блоки под заголовком (zoom), пока низ не встанет выше 1030 px.
FIT = ('<script>function fit(){var k=[].slice.call(document.body.children).filter(function(e){return !e.matches(".top,.tag,h1,h2,.plaque,.atom,script")});'
       'var z=1;function low(){var m=0;k.forEach(function(e){m=Math.max(m,e.getBoundingClientRect().bottom)});return m}'
       'while(low()>1030&&z>0.5){z-=0.04;k.forEach(function(e){e.style.zoom=z})}}'
       'fit();if(document.fonts)document.fonts.ready.then(fit);</script>')


def page(css: str, inner: str, cls: str) -> str:
    return f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><style>{css}</style></head><body class="{cls}">{inner}{FIT}</body></html>'


ATOM = ('<svg class="atom" viewBox="-200 -200 400 400">' + ''.join(
    f'<g transform="rotate({r})"><ellipse rx="180" ry="62" fill="none" stroke="#F3F0E8" stroke-width="10"/>'
    f'<circle cx="{c}" cy="{s}" r="15" fill="#F13A43"/></g>' for r, c, s in [(0, 120, 46), (60, -170, 20), (120, 60, -58)])
    + '<circle r="30" fill="#F13A43"/></svg>')


def blocks(body: str) -> tuple[str, str, str]:
    def grab(name):
        m = re.search(rf'(?:### {name}[^\n]*|\*\*{name}[^\n]*\*\*)\s*\n(.*?)(?=\n### |\n\*\*(?:Текст|Заметки|Раскрытие|Ответ|Формулы)[^\n]*\*\*\s*\n|\Z)', body, re.S)
        return m.group(1).strip() if m else ''
    return grab('Текст ученику'), grab('Раскрытие ответа'), grab('Заметки преподавателю')


def build(week: str) -> None:
    src = ROOT / 'course' / week / 'slides-outline.md'
    out_dir = ROOT / 'exports' / week
    out_dir.mkdir(parents=True, exist_ok=True)
    text = src.read_text(encoding='utf-8')
    readme = (ROOT / 'course' / week / 'README.md').read_text(encoding='utf-8')
    tm = re.search(r'^# Неделя\s*0?(\d+)[.:]\s*(.+)$', readme, re.M)
    week_no, topic = (tm.group(1), tm.group(2).strip()) if tm else (week[-2:].lstrip('0'), '')
    css = (HERE / 'slides.css').read_text(encoding='utf-8').replace('{{FONTS}}', (ROOT / 'content/brand/fonts').as_uri())
    label = f'<div class="top"><span><i></i>ХИМБИО ПРОСТО · НЕДЕЛЯ {week_no}</span><span>{{counter}}</span></div>'

    pages: list[tuple[str, str]] = []  # (html, notes)
    pages.append((page(css, f'<div class="top"><span><i></i>ХИМБИО ПРОСТО · ПОДГОТОВКА К ЕНТ</span><span></span></div>'
                       f'<div class="plaque">Неделя {week_no}</div><h1 class="cover-title">{html.escape(topic)}</h1>{ATOM}', 'dark cover'),
                  'Обложка недели.'))
    sections = list(re.finditer(r'^## Слайд (\d+)\.?\s*([^\n]*)', text, re.M))
    total = len(sections)
    for i, m in enumerate(sections):
        number, name = m.group(1), re.sub(r'\*\*|`', '', m.group(2)).strip()
        body = text[m.end():sections[i + 1].start() if i + 1 < total else len(text)].strip()
        visible, answer, notes = blocks(body)
        lines = [ln.strip() for ln in visible.splitlines() if ln.strip()]
        size = sum(len(ln) for ln in lines) + 45 * len(lines)
        density = 'dense' if size > 900 else 'medium' if size > 560 else ''
        main, calls = body_html(lines, dark=False)
        pages.append((page(css, label.replace('{counter}', f'{number} / {total}') + f'<h2><i></i><span>{inline(name)}</span></h2>'
                           f'<div class="body {density}">{main}</div><div class="callouts">{calls}</div>', 'light'),
                      (notes or body) + (f'\n\nОтвет (следующий слайд):\n{answer}' if answer else '')))
        if answer:
            alines = [ln.strip() for ln in answer.splitlines() if ln.strip()]
            asize = sum(len(ln) for ln in alines) + 45 * len(alines)
            adensity = 'dense' if asize > 700 else 'medium' if asize > 400 else ''
            amain, _ = body_html(alines, dark=True)
            pages.append((page(css, label.replace('{counter}', f'{number} / {total} · ответ') + f'<div class="plaque small">Ответ</div>'
                               f'<div class="answer-of">{inline(name)}</div><div class="body answer {adensity}">{amain}</div>', 'dark'),
                          'Показать после попытки учеников.\n\n' + answer))

    render(pages, out_dir / 'slides.pptx', out_dir / 'slides.pdf', keep=out_dir / 'slides-png' if '--keep-png' in sys.argv else None)
    print(f'Slides: {len(pages)} ({total} content + answers + cover) -> {out_dir / "slides.pptx"}, slides.pdf')


def render(pages: list[tuple[str, str]], pptx_path: Path, pdf_path: Path, keep: Path | None = None) -> None:
    """HTML-страницы 1920×1080 -> скриншоты Chrome -> PPTX (картинка + заметки) и PDF."""
    browser = chrome()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    with tempfile.TemporaryDirectory() as tmp:
        shots = []
        for k, (doc, _) in enumerate(pages, 1):
            h = Path(tmp) / f's{k:02d}.html'
            png = Path(tmp) / f's{k:02d}.png'
            h.write_text(doc, encoding='utf-8')
            subprocess.run([browser, '--headless=new', '--disable-gpu', '--hide-scrollbars', '--force-device-scale-factor=1',
                            '--window-size=1920,1080', f'--screenshot={png}', h.as_uri()], check=True, capture_output=True, timeout=120)
            shots.append(png)
        for png, (_, notes) in zip(shots, pages):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(str(png), 0, 0, deck.slide_width, deck.slide_height)
            slide.notes_slide.notes_text_frame.text = notes
        pptx_path.parent.mkdir(parents=True, exist_ok=True)
        deck.save(pptx_path)
        # PDF-версия: те же страницы подряд.
        all_html = Path(tmp) / 'all.html'
        all_html.write_text('<!doctype html><html><head><meta charset="utf-8"><style>@page{size:1920px 1080px;margin:0}'
                            'body{margin:0}img{display:block;width:1920px;height:1080px;break-after:page}</style></head><body>'
                            + ''.join(f'<img src="{p.as_uri()}">' for p in shots) + '</body></html>', encoding='utf-8')
        subprocess.run([browser, '--headless=new', '--disable-gpu', '--no-pdf-header-footer', f'--print-to-pdf={pdf_path}',
                        all_html.as_uri()], check=True, capture_output=True, timeout=180)
        if keep:
            keep.mkdir(parents=True, exist_ok=True)
            for p in shots:
                (keep / p.name).write_bytes(p.read_bytes())


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else 'week-01')
