"""Export the original Markdown lesson kit for review/import; never uploads files."""
from __future__ import annotations
import argparse
import re
from html import escape
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt as DocxPt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak

REPO = Path(__file__).resolve().parents[1]

def plain(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '').replace('<br>', '\n').strip()

def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or re.fullmatch(r'-{3,}', line):
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [plain(s) for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c.strip()) for c in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
            continue
        match = re.match(r'^(#{1,6})\s+(.+)', line)
        if match:
            yield 'heading', (len(match[1]), plain(match[2]))
            i += 1
            continue
        if line.startswith('```'):
            i += 1
            data = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                data.append(lines[i])
                i += 1
            yield 'paragraph', '\n'.join(data)
            i += 1
            continue
        if re.match(r'^[-*]\s+', line):
            yield 'paragraph', '• ' + plain(line[2:])
            i += 1
            continue
        data = [plain(line)]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||[-*] |```)', lines[i].strip()):
            data.append(plain(lines[i]))
            i += 1
        yield 'paragraph', ' '.join(data)

def page_break_title(title):
    match = re.match(r'Страница (\d+) из', title)
    return bool(match and int(match[1]) > 1)

def export_docx(text, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.6)
    doc.styles['Normal'].font.name = 'DejaVu Sans'
    doc.styles['Normal'].font.size = DocxPt(10.7)
    doc.styles['Normal'].paragraph_format.space_after = DocxPt(5)
    for kind, value in blocks(text):
        if kind == 'heading':
            level, title = value
            if page_break_title(title):
                doc.add_page_break()
            doc.add_heading(title, level=min(level, 3))
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=max(len(row) for row in value))
            table.style = 'Table Grid'
            for row in value:
                cells = table.add_row().cells
                for cell, content in zip(cells, row):
                    cell.text = content
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = DocxPt(9.5)
        else:
            doc.add_paragraph(value)
    section.footer.paragraphs[0].text = 'Химия просто • Черновик для проверки преподавателем'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    Document(output)  # Verify the written Office container opens.

def export_notes_pdf(text, output):
    font = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    bold = font.with_name('DejaVuSans-Bold.ttf')
    if not font.exists():  # Windows (Max's PC): Arial has Cyrillic.
        font, bold = Path('C:/Windows/Fonts/arial.ttf'), Path('C:/Windows/Fonts/arialbd.ttf')
    if not font.exists() or not bold.exists():
        raise RuntimeError('Install DejaVu Sans fonts before PDF export')
    pdfmetrics.registerFont(TTFont('DejaVu', str(font)))
    pdfmetrics.registerFont(TTFont('DejaVuBold', str(bold)))
    body = ParagraphStyle('body', fontName='DejaVu', fontSize=10.7, leading=14, spaceAfter=5)
    cell = ParagraphStyle('cell', parent=body, fontSize=9.4, leading=12, spaceAfter=2)
    heading = ParagraphStyle('heading', fontName='DejaVuBold', fontSize=12, leading=16, spaceBefore=7, spaceAfter=7)
    title_style = ParagraphStyle('title', parent=heading, fontSize=15, leading=19)
    width = A4[0] - 90
    story = []
    for kind, value in blocks(text):
        if kind == 'heading':
            level, title = value
            if page_break_title(title):
                story.append(PageBreak())
            story.append(Paragraph(escape(title), title_style if level == 1 else heading))
        elif kind == 'table':
            n = max(map(len, value))
            rows = [[Paragraph(escape(c).replace('\n', '<br/>'), cell) for c in row] for row in value]
            table = Table(rows, colWidths=[width/n]*n, repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.HexColor('#CBD5E1')),
                                       ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EFF6FF')),
                                       ('VALIGN',(0,0),(-1,-1),'TOP'),
                                       ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
            story += [table, Spacer(1,6)]
        else:
            story.append(Paragraph(escape(value).replace('\n','<br/>'), body))
    def footer(canvas, doc):
        canvas.setFont('DejaVu',8)
        canvas.drawString(45,24,'Химия просто • Черновик, ждёт проверки преподавателем')
        canvas.drawRightString(A4[0]-45,24,str(doc.page))
    output.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(output),pagesize=A4,leftMargin=45,rightMargin=45,topMargin=35,bottomMargin=38).build(story,onFirstPage=footer,onLaterPages=footer)

def export_slides(source, output):
    sections = list(re.finditer(r'^## Слайд (\d+)[^\n]*', source, re.M))
    if not sections:
        raise RuntimeError('Expected headings: ## Слайд N. Title')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, match in enumerate(sections):
        body = source[match.end():sections[index+1].start() if index+1<len(sections) else len(source)].strip()
        # Visible text/teacher notes are deliberately separate in the Markdown source.
        visible_match = re.search(r'(?:### Текст ученику|\*\*Текст ученику[^\n]*\*\*)\s*\n(.*?)(?=\n(?:### |\*\*(?:Заметки|Ответ|Раскрытие|Формулы))|\Z)',body,re.S)
        if not visible_match:
            raise RuntimeError(f'Missing visible text block on slide {match[1]}')
        visible = plain(visible_match[1])
        visible_lines=[]
        for line in visible.splitlines():
            line=line.strip()
            if not line:continue
            if line.startswith('|'):
                cells=[c.strip() for c in line.strip('|').split('|')]
                if all(re.fullmatch(r':?-+:?',c) for c in cells):continue
                line='  ·  '.join(c for c in cells if c)
            visible_lines.append(line)
        estimated_lines=sum(max(1,(len(line)+69)//70) for line in visible_lines)
        body_size=22 if estimated_lines<=12 else 20 if estimated_lines<=16 else 18
        if estimated_lines>20:raise RuntimeError(f'Slide {match[1]} too dense; revise visible text before export')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(248,250,252)
        title = slide.shapes.add_textbox(Inches(.7), Inches(.4), Inches(12), Inches(.9)).text_frame
        p = title.paragraphs[0]
        p.text = plain(match[0].lstrip('# ').strip())
        p.font.size, p.font.bold = Pt(28), True
        p.font.color.rgb = RGBColor(15,23,42)
        frame = slide.shapes.add_textbox(Inches(.8),Inches(1.55),Inches(11.8),Inches(5.15)).text_frame
        frame.word_wrap = True
        for j, line in enumerate(visible_lines):
            if not line.strip():continue
            p = frame.paragraphs[0] if not frame.text else frame.add_paragraph()
            p.text = line.lstrip('- ').strip()
            p.font.name = 'DejaVu Sans'
            p.font.size = Pt(body_size)
            p.font.color.rgb = RGBColor(30,41,59)
            p.space_after = Pt(3)
        slide.notes_slide.notes_text_frame.text = body
        foot = slide.shapes.add_textbox(Inches(.8),Inches(7),Inches(11.8), Inches(.3)).text_frame
        foot.text = f'Черновик для проверки преподавателем • {match[1]}'
        foot.paragraphs[0].font.size = Pt(12)
    output.parent.mkdir(parents=True,exist_ok=True)
    deck.save(output)
    restored=Presentation(output)
    if len(restored.slides)!=len(sections):raise RuntimeError('Incomplete presentation export')
    print(f'PPTX: {len(restored.slides)} slides, teacher notes retained')

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--week',default='week-01')
    args=parser.parse_args()
    source=(REPO/'course'/args.week).resolve()
    if source.parent != (REPO/'course').resolve():raise ValueError('Expected course/week-NN')
    output=REPO/'exports'/args.week
    student={'notes-student','pre-work','homework'}
    for name in ['notes-student','pre-work','homework','lesson-1-plan','lesson-2-plan','homework-solutions','teacher-notes','README']:
        text=(source/f'{name}.md').read_text(encoding='utf-8')
        export_docx(text,output/('student' if name in student else 'teacher')/f'{name}.docx')
    quiz=(source/'quiz.md').read_text(encoding='utf-8')
    boundary=re.search(r'^## (?:Ключ|Для преподавателя|Преподавател|Только преподавателю)',quiz,re.M)
    if not boundary:raise RuntimeError('Expected separate teacher section in quiz.md')
    export_docx(quiz[:boundary.start()],output/'student/quiz.docx')
    export_docx(quiz,output/'teacher/quiz-with-key.docx')
    export_notes_pdf((source/'notes-student.md').read_text(encoding='utf-8'),output/'student/notes-student.pdf')
    export_slides((source/'slides-outline.md').read_text(encoding='utf-8'),output/'slides.pptx')
    print(f'Local review/import files: {output}; no Google upload was performed')

if __name__=='__main__':main()
